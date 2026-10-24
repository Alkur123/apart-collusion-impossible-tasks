"""Pour report/REPORT.md (after fill_report.py) into the official Apart template, then export PDF via Word.

  python report/build_report.py            # -> report/report.docx, report/report.pdf
Markdown subset: '## ', '### ', paragraphs, '- ' bullets, '1. ' items, pipe tables,
'**bold**', '*italic*', '`code`', and a line '{{FIG:path|caption}}' for a figure.
"""
import copy
import pathlib
import re
import subprocess
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

HERE = pathlib.Path(__file__).resolve().parent
TEMPLATE = HERE / "template.docx"
SRC = HERE / "REPORT_FILLED.md"
FONT = "Old Standard TT"
BODY_PT = 10.5


def set_t(t_el, text):
    t_el.text = text
    t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def header(doc, title, author, affil, abstract):
    body = doc.element.body
    ts = [t for t in body.iter() if t.tag.endswith("}t")]
    for t in ts:
        if t.text == "PROJECT TITLE":
            set_t(t, title)
    # authors: the template's 6-slot author sub-table collapses around one filled slot,
    # so replace that sub-table with two centred lines (name, affiliation).
    first = next(t for t in ts if t.text == "Author name 1")
    grid = first
    while not grid.tag.endswith("}tbl"):
        grid = grid.getparent()
    parent = grid.getparent()
    idx = list(parent).index(grid)
    parent.remove(grid)
    tmp = docx.Document()
    for k, (text, size) in enumerate([(author, 11), (affil, 10)]):
        p = tmp.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(size)
        parent.insert(idx + k, copy.deepcopy(p._p))
    for t in ts:
        if t.text and t.text.startswith("Summarize your project in 150"):
            set_t(t, abstract)
            # drop the remaining guidance runs in the same paragraph
            p = t.getparent().getparent()
            for r in list(p):
                if r.tag.endswith("}r") and r is not t.getparent():
                    for tt in r.iter():
                        if tt.tag.endswith("}t") and tt is not t:
                            tt.text = ""


def strip_guidance(doc):
    """Remove everything after the header table (info box + guidance), keep sectPr."""
    body = doc.element.body
    kids = list(body)
    first_tbl = next(i for i, e in enumerate(kids) if e.tag.endswith("}tbl"))
    for e in kids[first_tbl + 1:]:
        if not e.tag.endswith("}sectPr"):
            body.remove(e)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(p, text, size=BODY_PT, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(size - 1)
            continue
        elif part.startswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        r.font.name = FONT
        r.font.size = Pt(size)
        if italic:
            r.italic = True


def para(doc, text, style="normal", size=BODY_PT, indent=None, align=None, space_after=4):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, size)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.08
    if indent is not None:
        pf.left_indent = Inches(indent)
        pf.first_line_indent = Inches(-0.17)
    if align:
        p.alignment = align
    return p


def _borders(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "BFBFBF")
        b.append(e)
    tblPr.append(b)


def table(doc, rows):
    hdr = [c.strip() for c in rows[0].strip("|").split("|")]
    body = [[c.strip() for c in r.strip("|").split("|")] for r in rows[2:]]
    t = doc.add_table(rows=1 + len(body), cols=len(hdr))
    t.autofit = False
    _borders(t)
    total = 6.5
    first = 2.3 if len(hdr) > 3 else total / len(hdr)
    rest = (total - first) / max(1, len(hdr) - 1)
    for j in range(len(hdr)):
        for cell in t.columns[j].cells:
            cell.width = Inches(first if j == 0 else rest)
    for i, r in enumerate([hdr] + body):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_runs(p, f"**{v}**" if i == 0 else v, 8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    md = SRC.read_text(encoding="utf-8").splitlines()
    meta = {m.group(1): m.group(2) for l in md if (m := re.match(r"% (\w+): (.*)", l))}
    doc = docx.Document(str(TEMPLATE))
    header(doc, meta["TITLE"], meta["AUTHOR"], meta["AFFILIATION"], meta["ABSTRACT"])
    strip_guidance(doc)
    lines = [l for l in md if not l.startswith("% ")]
    i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("## "):
            h = doc.add_paragraph(style="Heading 2"); add_runs(h, l[3:], 13)
            h.paragraph_format.space_before = Pt(8); h.paragraph_format.space_after = Pt(3)
        elif l.startswith("### "):
            h = doc.add_paragraph(style="Heading 3"); add_runs(h, l[4:], 11.5)
            h.paragraph_format.space_before = Pt(4); h.paragraph_format.space_after = Pt(2)
        elif l.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i]); i += 1
            table(doc, rows)
            continue
        elif l.startswith("{{FIG:"):
            path, cap = l[6:-2].split("|", 1)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(HERE / path), width=Inches(6.4))
            para(doc, cap, size=9, space_after=8)
        elif l.startswith("- "):
            para(doc, "•  " + l[2:], indent=0.25, space_after=2)
        elif re.match(r"\d+\. ", l):
            para(doc, l, indent=0.25, space_after=2)
        elif l.startswith("[") and re.match(r"\[\d+\]", l):
            para(doc, l, size=9, indent=0.3, space_after=1)
        elif l.strip():
            para(doc, l)
        i += 1
    out = HERE / "report.docx"
    doc.save(str(out))
    return out


def to_pdf(docx_path):
    pdf = docx_path.with_suffix(".pdf")
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{docx_path}'); $d.SaveAs([ref]'{pdf}', [ref]17); $d.Close(); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    return pdf


if __name__ == "__main__":
    d = build()
    print("docx", d)
    if "--no-pdf" not in sys.argv:
        print("pdf", to_pdf(d))
